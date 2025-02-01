import openpyxl
import re

def reemplazar_rutas_excel(archivo, ruta_original, ruta_nueva, archivo_salida):
    # Cargar el archivo de Excel
    wb = openpyxl.load_workbook(archivo, keep_links=True)
    
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        
        for fila in ws.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and ruta_original in celda.value:
                    nueva_formula = celda.value.replace(ruta_original, ruta_nueva)
                    celda.value = nueva_formula
                    print(f"Modificada celda {celda.coordinate}: {nueva_formula}")
    
    # Guardar el archivo modificado
    wb.save(archivo_salida)
    print(f"Archivo guardado como {archivo_salida}")

# Parámetros
archivo_excel = "Reporte_Dip_SinInc_vf_2.xlsx"  # Reemplaza con la ruta de tu archivo
ruta_vieja = "https://d.docs.live.net/Users/Ivan/Desktop/Ivan/Laboral/EstrategiaSur/GitProjects/Proyeccion-electoral-congreso/Resultados"
ruta_nueva = "C:\\Users\\pablo\\OneDrive\\Documents\\Trabajo\\EstrategiaSur\\Proyeccion-electoral-congreso\\Resultados_correcciones"
archivo_salida = "Reporte_Dip_SinInc_vf_2_corregido.xlsx"

# Ejecutar la función
reemplazar_rutas_excel(archivo_excel, ruta_vieja, ruta_nueva, archivo_salida)